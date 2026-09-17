"""Read-only, CPU-only extraction for docs/audits/exr6_evidence.md.

Requires existing caches and the raw release files; never compiles or saves a
connectome. External source downloads go only to the requested output folder.
ExR2_1/ExR2_2 are grouped for this family comparison, not renamed in the model.
"""
from pathlib import Path
import argparse
import hashlib
import json
import subprocess
import urllib.request

import numpy as np
import pandas as pd
import pyarrow as pa
import pyarrow.compute as pc
import pyarrow.feather as pf
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
FAMILY = [f"ExR{i}" for i in range(1, 9)] + ["ER6", "ER4m"]
NTS = ["acetylcholine", "dopamine", "gaba", "glutamate", "histamine", "octopamine", "serotonin"]
GENES = ["GluClalpha", "mGluR", "Rdl", "GABA-B-R1", "GABA-B-R2", "GABA-B-R3"]
WOLFF = "https://cdn.elifesciences.org/articles/104764/"
WOLFF_SHA256 = {
    2: "9ab5996b4f87abfe625f2e2789ddfa900c13489f4231b4b2be12c9e18209e2d1",
    9: "d6a81bd862dc2de8fdc34cd38d886423a1b36bb004374545b339e314a7d5fb40",
}


def digest(path, algorithm="sha256"):
    h = hashlib.new(algorithm)
    with Path(path).open("rb") as f:
        for chunk in iter(lambda: f.read(4 * 1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def family(names):
    return names.str.strip().replace({"ExR2_1": "ExR2", "ExR2_2": "ExR2"})


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--cache", type=Path, default=ROOT / "cache")
    parser.add_argument("--male-raw", type=Path, default=Path("D:/Datasets/male-cns-connectome-v1.0/flat-connectome"))
    parser.add_argument("--female-raw", type=Path, default=Path("D:/Datasets/flywire"))
    parser.add_argument("--external", type=Path, default=ROOT / "data/external")
    parser.add_argument("--out", type=Path, default=ROOT / "out/exr6_evidence")
    args = parser.parse_args()
    args.out.mkdir(parents=True, exist_ok=True)
    sources = args.out / "sources"
    sources.mkdir(exist_ok=True)
    inputs, cache_before, cache_manifest = {}, {}, {}

    def record(path, key=None, url=None):
        path = Path(path)
        inputs[key or path.name] = dict(path=str(path.resolve()), bytes=path.stat().st_size,
                                       sha256=digest(path), **({"url": url} if url else {}))
        return path

    def emit(frame, name):
        frame.to_csv(args.out / (name + ".csv"), index=False, lineterminator="\n")

    cells, summaries = {}, []
    for dataset in ("malecns", "fafb", "banc"):
        cache = args.cache / ("" if dataset == "malecns" else dataset)
        # Fingerprint all top-level cache files, including weights and sign-zero counts.
        for p in sorted(cache.iterdir()):
            if p.is_file():
                cache_before[str(p.resolve())] = digest(p, "md5")
        manifest = cache / "manifest.json"
        cache_manifest[dataset] = json.loads(manifest.read_text()) if manifest.exists() else {"dataset": dataset, "release": "v1.0", "legacy_cache": True}
        n = pd.read_parquet(record(cache / "neurons.parquet", dataset + "_neurons"))
        n["family"] = family(n.type)
        n = n.loc[n.family.isin(FAMILY), ["bodyId", "type", "family", "nt", "sign", "somaSide"]].copy()
        cells[dataset] = n
        for ty in FAMILY:
            s = n.loc[n.family.eq(ty)]
            summaries.append(dict(dataset=dataset, family=ty, n=len(s),
                                  compiled_labels=json.dumps(s.nt.value_counts().to_dict(), sort_keys=True)))

    male = pd.read_feather(record(args.male_raw / "body-neurotransmitters-male-cns-v1.0.feather"))
    male = cells["malecns"].merge(male, left_on="bodyId", right_on="body", validate="one_to_one")
    emit(male, "male_cells")
    # The argmax fractions are votes over all seven classes, not mean softmax
    # probabilities and not the confidence field in the release's body table.
    pa.set_cpu_count(2)
    cols = ["body"] + [f"nt_{nt}_prob" for nt in NTS]
    table = pf.read_table(record(args.male_raw / "tbar-neurotransmitters-male-cns-v1.0.feather"),
                          columns=cols, memory_map=True)
    table = table.filter(pc.is_in(table["body"], value_set=pa.array(male.bodyId.to_numpy(), type=pa.int64())))
    tbar = table.to_pandas()
    probs = tbar[cols[1:]].to_numpy()
    if not np.isfinite(probs).all():
        raise ValueError("Non-finite T-bar probabilities; refusing ambiguous argmax")
    votes = probs.argmax(axis=1)
    ties = (probs == probs.max(axis=1, keepdims=True)).sum(axis=1) > 1
    counts = pd.DataFrame({"bodyId": tbar.body.to_numpy(), "tied_argmax": ties.astype(int)})
    for i, nt in enumerate(NTS):
        counts[nt] = (votes == i).astype(int)
        counts["prob_sum_" + nt] = probs[:, i].astype(np.float64)
    counts["n_tbars"] = 1
    per_body = counts.groupby("bodyId", as_index=False).sum().merge(male[["bodyId", "family"]], validate="one_to_one")
    per_type = per_body.drop(columns="bodyId").groupby("family", as_index=False).sum()
    for frame, name in ((per_body, "male_tbars_by_body"), (per_type, "male_tbars_by_type")):
        for nt in NTS:
            frame["fraction_" + nt] = frame[nt] / frame.n_tbars
            frame["mean_probability_" + nt] = frame["prob_sum_" + nt] / frame.n_tbars
        emit(frame, name)
    assert np.array_equal(per_body.n_tbars.to_numpy(), male.set_index("bodyId").total_nt_predictions.reindex(per_body.bodyId).to_numpy()), "T-bar counts disagree with body release"

    fafb = pd.read_csv(record(args.female_raw / "Female Adult Fly Brain v783/neurons.csv.gz", "fafb_raw_neurons"), low_memory=False)
    fafb = cells["fafb"].merge(fafb, left_on="bodyId", right_on="root_id", validate="one_to_one")
    known_path = record(args.external / "typing/schlegel2024_Supplemental_file1_neuron_annotations.tsv")
    known = pd.read_csv(known_path, sep="\t", low_memory=False)
    fields = ["root_id", "cell_type", "hemibrain_type", "top_nt", "top_nt_conf", "known_nt", "known_nt_source"]
    fafb = fafb.merge(known[fields], on="root_id", validate="one_to_one")
    emit(fafb, "fafb_cells")
    banc = pd.read_csv(record(args.female_raw / "BANC v888/neurons.csv.gz", "banc_raw_neurons"), low_memory=False)
    fields = ["Root ID", "Primary Cell Type", "Alternative Cell Type(s)", "Predicted NT type",
              "Predicted NT confidence", "Verified NT type", "Verified Neuropeptide"]
    banc = cells["banc"].merge(banc[fields], left_on="bodyId", right_on="Root ID", validate="one_to_one")
    emit(banc, "banc_cells")
    for d, raw in (("malecns", male), ("fafb", fafb), ("banc", banc)):
        assert len(raw) == len(cells[d]), f"Missing raw rows for {d}"
    emit(pd.DataFrame(summaries), "family_summary")

    # Store the named primary worksheets. Their hashes belong in report.json;
    # no live web request or annotation substitution is part of connectome.load().
    for figure in (2, 9):
        filename = f"elife-104764-fig{figure}-data1-v1.xlsx"
        p = sources / filename
        if not p.exists():
            with urllib.request.urlopen(WOLFF + filename) as response:
                p.write_bytes(response.read())
        record(p, url=WOLFF + filename)
        if digest(p) != WOLFF_SHA256[figure]:
            raise ValueError(f"Wolff figure {figure} source bytes differ from the audited version")
        t = pd.read_excel(p).rename(columns=lambda c: str(c).strip())
        type_col = "Cell type"
        if type_col not in t:
            raise ValueError(f"Unexpected Wolff figure {figure} sheet schema: {list(t)}")
        t["family"] = family(t[type_col])
        emit(t[t.family.isin(FAMILY + ["EPG", "PEN_a(PEN1)", "PEN_b(PEN2)"])], f"wolff_fig{figure}")
        if figure == 9:
            rows = []
            for row in load_workbook(p).active.iter_rows(max_col=6):
                if str(row[1].value).strip() in ("ExR6", "ER6"):
                    rows.append(dict(type=str(row[1].value).strip(), cell=row[3].coordinate,
                                     marker=row[3].value, bold=row[3].font.b, italic=row[3].font.i,
                                     probe_sets=row[4].value, peptide_result=row[5].value))
            emit(pd.DataFrame(rows), "wolff_focus_format")

    targets = FAMILY + ["EPG", "PEN_a(PEN1)", "PEN_b(PEN2)"]
    maps = []
    for p in sorted((ROOT / "flyverse/data").glob("type_map_*.csv")):
        t = pd.read_csv(record(p), comment="#")
        if "malecns_type" in t:
            t = t[t.malecns_type.isin(targets)].copy()
            t.insert(0, "map_file", p.name)
            maps.append(t)
    emit(pd.concat(maps, ignore_index=True), "type_maps")
    nern = []
    for name in ("nern2025_Nern-et-al_SuppTable01_Cell-types-and-counts.xlsx",
                 "nern2025_Nern-et-al_SuppTable05_Neurotransmitter_validation.xlsx"):
        p = record(args.external / "typing" / name)
        for sheet, t in pd.read_excel(p, sheet_name=None, header=None).items():
            matched = t.astype(str).apply(lambda c: c.str.strip().isin(FAMILY)).any(axis=1)
            nern.append(dict(file=name, sheet=sheet, exact_family_rows=int(matched.sum())))
    emit(pd.DataFrame(nern), "nern_family_coverage")
    receptors = pd.read_csv(record(ROOT / "flyverse/data/receptors_by_type.csv"), comment="#")
    emit(receptors[receptors.malecns_type.isin(targets)], "receptor_rows")
    expr = pd.read_csv(record(ROOT / "flyverse/data/expression_davis2020.csv"), comment="#")
    emit(expr[expr.source_name.isin(["PB_2", "PB_3"])][["source_name", "qc", "n_samples_pass", "n_samples_suboptimal"] +
          [g + suffix for g in GENES for suffix in ("_tpm", "_p_on")]], "target_expression")

    cache_after = {p: digest(p, "md5") for p in cache_before}
    assert cache_after == cache_before, "A cache file changed during extraction"
    outputs = {p.name: digest(p) for p in sorted(args.out.glob("*.csv"))}
    report = dict(schema=1, base_commit="8854691", code_commit=subprocess.check_output(
        ["git", "rev-parse", "HEAD"], cwd=ROOT, text=True).strip(),
        script_sha256=digest(__file__), inputs=inputs, output_sha256=outputs,
        cache_manifest=cache_manifest, cache_md5_before=cache_before, cache_md5_after=cache_after,
        cache_unchanged=True, tbar_argmax_ties=int(ties.sum()),
        tbar_count_matches_body_table=True,
        family_grouping={"ExR2": ["ExR2", "ExR2_1", "ExR2_2"]},
        note="Read-only CPU extraction; no model, cache, override, or receptor change.")
    (args.out / "report.json").write_text(json.dumps(report, indent=2, ensure_ascii=True) + "\n", encoding="ascii")
    print(pd.DataFrame(summaries).to_string(index=False))
    print(f"Saved {len(outputs)} CSVs and report.json; cache bytes unchanged.")


if __name__ == "__main__":
    main()
